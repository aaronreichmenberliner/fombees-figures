"""Read ESM_BM per scenario from the draft 8 workbook.

The five category rows are summed rather than read from each sheet's own total,
because two sheets' totals are known to be wrong (see check_esmbm_rollup).

Label handling matters more than it looks. The sheets do not spell the five
categories identically - FLOWN writes "Waste outputs" where the others write
"Waste process outputs", and several sheets leave a trailing space on
"Process Inputs ". Matching only the canonical spelling silently drops a row,
which is exactly how 124 kgeq of FLOWN waste went missing through five
published figures.
"""
from __future__ import annotations

import openpyxl

from fombees_paths import workbook

CATS = ["System", "Process Operations", "Process Inputs",
        "Waste process outputs", "Useful process outputs"]
ALIAS = {
    "waste outputs": "Waste process outputs",
    "useful outputs": "Useful process outputs",
    "inputs": "Process Inputs",
    "process operations": "Process Operations",
}
SCEN = ["FLOWN", "TRANSGENIC-LETTUCE", "TRANSGENIC-TOBACCO", "GENE GUN",
        "AGRO", "VIRAL"]


def _canon(label: str | None) -> str | None:
    if not isinstance(label, str):
        return None
    s = " ".join(label.split())          # collapse stray whitespace
    for cat in CATS:
        if s.lower() == cat.lower():
            return cat
    return ALIAS.get(s.lower())


def blocks(sheet_v) -> list[dict]:
    """Every 'Category' block on a sheet, with its header and category values."""
    out, cur, hdr = [], None, ""
    for r in range(1, sheet_v.max_row + 1):
        a = sheet_v.cell(r, 1).value
        a = str(a).strip() if a else ""
        if a.startswith(("Equivalent System Mass", "Complete ESM")):
            hdr = a
        if a == "Category":
            cur = {"_hdr": hdr, "_row": r}
            out.append(cur)
        cat = _canon(a)
        if cur is not None and cat and cat not in cur:
            v = sheet_v.cell(r, 2).value
            cur[cat] = float(v) if isinstance(v, (int, float)) else 0.0
    return out


# Which roll-up row is a scenario's authoritative total. Matched on the label
# rather than a fixed address: deleting a row shifts every cell below it, and
# hardcoded addresses silently start reading the wrong thing.
TOTAL_LABEL = {
    "FLOWN": "ESMBM (kgeq)",
    "TRANSGENIC-LETTUCE": "ESMBM (kgeq)",
    "TRANSGENIC-TOBACCO": "ESMBM (kgeq)",
    "GENE GUN": "Total Overall ESM (kgeq)",
    "AGRO": "Total Overall ESMBM (kgeq)",
    "VIRAL": "Total Overall ESMBM (kgeq)",
}


def total_cells() -> dict[str, tuple[str, str]]:
    """scenario -> (sheet, cell) holding its own ESM_BM total, found by label."""
    V = openpyxl.load_workbook(workbook(), data_only=True)
    out = {}
    for scen, label in TOTAL_LABEL.items():
        sheet = V[scen]
        for r in range(1, sheet.max_row + 1):
            a = sheet.cell(r, 1).value
            if isinstance(a, str) and a.strip() == label:
                out[scen] = (scen, f"B{r}")
                break
        else:
            raise LookupError(f"{scen}: no row labelled {label!r}")
    return out


def read() -> dict[str, dict[str, float]]:
    V = openpyxl.load_workbook(workbook(), data_only=True)
    res = {}
    for s in SCEN:
        bs = blocks(V[s])
        combined = [b for b in bs if b["_hdr"].startswith("Complete ESM")]
        use = combined or [b for b in bs if "System" in b]
        if not use:
            raise LookupError(f"{s}: no category block found")
        p = {c: sum(b.get(c, 0.0) for b in use) for c in CATS}
        p["ESM_BM"] = (p["System"] + p["Process Operations"]
                       + p["Process Inputs"] + p["Waste process outputs"]
                       - p["Useful process outputs"])
        p["GROSS"] = (p["System"] + p["Process Operations"]
                      + p["Process Inputs"] + p["Waste process outputs"])
        res[s] = p
    return res


if __name__ == "__main__":
    r = read()
    print(f"{'':<20}{'System':>10}{'ProcOps':>10}{'Inputs':>10}"
          f"{'Waste':>9}{'Useful':>10}{'NET':>12}")
    for s, v in sorted(r.items(), key=lambda kv: kv[1]["ESM_BM"]):
        print(f"{s:<20}{v['System']:>10,.1f}{v['Process Operations']:>10,.1f}"
              f"{v['Process Inputs']:>10,.1f}{v['Waste process outputs']:>9,.1f}"
              f"{v['Useful process outputs']:>10,.1f}{v['ESM_BM']:>12,.1f}")
