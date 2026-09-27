"""Read the per-component figure-of-merit tables out of the draft 8 workbook.

Each scenario sheet carries a "FOM Graphs" block whose two header rows define
four metric groups. The block is located by its own label rather than by a fixed
offset, because the blocks sit at different rows on every sheet.

Layout, identical on all six sheets:
    Mass       B C D E      System, Input, Output, Waste
    Volume     F G H I      System, Input, Output, Waste
    Power      J K L M      System, Input, Output, Waste
    Crew time  N O P Q R    System, Process Operations, Input, Output, Waste
"""
from __future__ import annotations

import openpyxl

from fombees_paths import workbook

SCEN = ["FLOWN", "TRANSGENIC-LETTUCE", "TRANSGENIC-TOBACCO", "GENE GUN",
        "AGRO", "VIRAL"]

# metric -> {flow: column index}
COLUMNS = {
    "mass":      {"System": 2, "Input": 3, "Output": 4, "Waste": 5},
    "volume":    {"System": 6, "Input": 7, "Output": 8, "Waste": 9},
    "power":     {"System": 10, "Input": 11, "Output": 12, "Waste": 13},
    "crew_time": {"System": 14, "Process Operations": 15, "Input": 16,
                  "Output": 17, "Waste": 18},
}
CREDIT = "Output"

# The sheets spell several components more than one way.
RENAME = {
    "Biomass Waste": "Biomass waste",
    "Harvesting collection container": "Harvesting container",
    "N. benthamiana seeds": "N benthamiana seeds",
    "Plant nutrients": "Plant growth nutrients",
    "Cartrdige extractor": "Cartridge extractor",   # typo in the source
}

GROUP = {
    "Biomass & product": [
        "PTH-Fc", "Plant biomass", "Biomass waste", "Calories",
        "Macronutrients", "Micronutrients"],
    "Cold storage": [
        "Refrigerator/Freezer"],
    "Plant growth & harvest": [
        "Plant growth chamber", "Lighting panels", "Harvesting tools",
        "Harvesting container", "Lettuce seeds", "N benthamiana seeds"],
    "Processing & inoculation": [
        "Gene Gun", "BSC", "Shaking incubator", "Autoclave", "Centrifuge",
        "Vacuum pump", "Vacuum chamber", "Spectrophotometer", "Scale",
        "Battery", "Gas cylinder", "Gas regulator", "Connector tubing",
        "Bullets", "Bullet holder vial", "Cartridge extractor",
        "Dessicant pack", "Viral stock", "Carborundum"],
    "Labware & packaging": [
        "Flask", "Cuvette", "Bottle", "Beaker", "Culture tubes",
        "Centrifuge tubes", "Agro stock tubes", "Mixing tube",
        "Eye protection", "Ear protection", "Injection Pens (packaged)",
        "Used Injection Pens", "Waste Packaging"],
    "Gases (CO₂, O₂)": [
        "Carbon dioxide", "Oxygen"],
    "Water": [
        "Water"],
    "Nutrients & media": [
        "Plant growth nutrients", "LB Media mix", "Kanamycin", "Rifampicin",
        "Tetracycline", "MES", "MgCl2", "Acetosyringone", "Silwet L-77"],
}
OF_GROUP = {c: g for g, items in GROUP.items() for c in items}


def _block_row(sheet) -> int:
    for r in range(1, sheet.max_row + 1):
        v = sheet.cell(r, 1).value
        if isinstance(v, str) and v.strip() == "FOM Graphs":
            return r
    raise LookupError(f"no 'FOM Graphs' block on {sheet.title}")


def _verify_header(sheet, hdr: int) -> None:
    """Fail if the columns are not where COLUMNS says they are."""
    for metric, flows in COLUMNS.items():
        for flow, col in flows.items():
            got = sheet.cell(hdr + 1, col).value
            if got is None or str(got).strip() != flow:
                raise AssertionError(
                    f"{sheet.title}: expected {flow!r} at column {col} of the "
                    f"{metric} group, found {got!r}. The block layout has "
                    f"changed and the column map must be updated.")


def rows() -> list[dict]:
    """One record per (scenario, component, metric, flow) with a non-zero value."""
    V = openpyxl.load_workbook(workbook(), data_only=True)
    out, unmapped = [], set()
    for scen in SCEN:
        sh = V[scen]
        start = _block_row(sh)
        _verify_header(sh, start + 1)
        r = start + 3
        while r <= sh.max_row:
            name = sh.cell(r, 1).value
            if name is None or not str(name).strip():
                break
            comp = RENAME.get(str(name).strip(), str(name).strip())
            group = OF_GROUP.get(comp)
            if group is None:
                unmapped.add(f"{scen}!A{r} {comp!r}")
            for metric, flows in COLUMNS.items():
                for flow, col in flows.items():
                    v = sh.cell(r, col).value
                    if isinstance(v, (int, float)) and v:
                        out.append({"scenario": scen, "component": comp,
                                    "group": group, "metric": metric,
                                    "flow": flow, "value": float(v)})
            r += 1
    if unmapped:
        raise AssertionError(
            "components with no resource group - add them to GROUP:\n  "
            + "\n  ".join(sorted(unmapped)))
    return out


if __name__ == "__main__":
    rs = rows()
    comps = {r["component"] for r in rs}
    print(f"{len(rs)} values, {len(comps)} components, "
          f"{len({r['group'] for r in rs})} groups")
    for scen in SCEN:
        sub = [r for r in rs if r["scenario"] == scen]
        print(f"  {scen:<20} {len(sub):>4} values, "
              f"{len({r['component'] for r in sub}):>3} components")
